# -*- coding: utf-8 -*-
"""
कार्य-वृत्त — Flask बैकएंड
===========================
- /form            → सार्वजनिक हिंदी फॉर्म — प्रथम भाग + द्वितीय भाग (कोई लॉगिन नहीं)
- /submit          → उत्तर सहेजना (POST, योग स्वतः)
- /dhanyavaad      → धन्यवाद पृष्ठ
- /admin           → एडमिन डैशबोर्ड (टेबल + इन्फोग्राफिक्स + Excel + PDF)
- /admin/export    → Excel (.xlsx) डाउनलोड — हमेशा ताज़ा, single combined
- /admin/pdf1      → PDF-1 (प्रथम भाग) — सभी उत्तरों की table
- /admin/pdf2      → PDF-2 (द्वितीय भाग) — सभी उत्तरों की table
- /admin/clear     → सारा डेटा साफ़ करना
- /admin/api/data  → डैशबोर्ड द्वारा 10 सेकंड में ऑटो-रिफ्रेश वाला JSON

माह/वर्ष हर माह अपने आप बदलता है (survey_config.get_form_title)।
"""
import io
import os
import json
import time
import hashlib
import sqlite3
import hmac
import secrets
from datetime import datetime, date, timedelta

from flask import (
    Flask, request, render_template, redirect, url_for,
    session, jsonify, send_file, flash, abort
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fpdf import FPDF
from fpdf.fonts import FontFace

from survey_config import (
    SECTIONS, FORM_DESC, THANK_YOU_MSG,
    ADMIN_PASSWORD, SITE_TITLE, ORG_NAME, ORG_TAGLINE, ORG_SUB, EXPECTED_PROVINCES,
    PART_ORDER, PART_SUBTITLES, PART_SHORT, OTHER_PROV_LABEL,
    get_form_title, get_instruction, get_month_label, current_month_year,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "data", "responses.db")
FONTS_DIR = os.path.join(BASE_DIR, "static", "fonts")
LOGO_EXISTS = os.path.exists(os.path.join(BASE_DIR, "static", "logo.png"))

app = Flask(__name__)


def _load_secret_key():
    """SECRET_KEY: पर्यावरण-चर से, वरना स्थायी फ़ाइल से (रीस्टार्ट पर भी वही रहे)"""
    env = os.environ.get("SECRET_KEY")
    if env:
        return env
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    path = os.path.join(os.path.dirname(DB_PATH), "secret_key.txt")
    try:
        with open(path) as fh:
            k = fh.read().strip()
            if k:
                return k
    except FileNotFoundError:
        pass
    k = secrets.token_hex(32)
    with open(path, "w") as fh:
        fh.write(k)
    return k


app.secret_key = _load_secret_key()


# ---------------------------------------------------------------- सपाट फ़ील्ड सूची
def flat_fields():
    """सभी फ़ील्ड एक सपाट सूची में — क्रम: प्रथम भाग → द्वितीय भाग"""
    out = []
    for sec in SECTIONS:
        for item in sec.get("items", []):
            for f in item.get("fields", []):
                out.append({
                    "id": f["id"],
                    "label": f.get("label", ""),
                    "short": f.get("short", ""),
                    "type": f.get("type", "text"),
                    "part": sec.get("part", ""),
                    "section": sec.get("title", ""),
                    "grp": item.get("grp", ""),
                    "item_label": item.get("label", ""),
                    "required": bool(f.get("required", True)),
                    "full": bool(f.get("full", False)),
                    "readonly": bool(f.get("readonly", False)),
                    "calc": f.get("calc"),
                    "options": f.get("options") or [],
                    "show_if": f.get("show_if"),
                    "show_if_value": f.get("show_if_value"),
                    "autofill_from": f.get("autofill_from"),
                    "placeholder": f.get("placeholder", ""),
                })
    return out


FLAT = flat_fields()
BY_ID = {f["id"]: f for f in FLAT}


def col_header(f):
    """Excel हेडर — 'प्रथम भाग · ख.1 — जिलों की वर्तमान संख्या' जैसा"""
    if f["grp"]:
        return f"{f['part']} · {f['grp']} — {f['label']}"
    return f"{f['part']} · {f['label']}"


def to_num(v):
    """स्ट्रिंग → संख्या (जितना हो सके)"""
    v = (v or "").strip()
    if not v:
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except ValueError:
        return None


def display_prov(data):
    """दिखाने वाला प्रांत नाम — 'अन्य' चुना हो तो लिखा हुआ नाम।"""
    prov = (data.get("p1_prov") or "").strip()
    if prov == OTHER_PROV_LABEL:
        return (data.get("p1_prov_other") or "").strip()
    return prov


# ---------------------------------------------------------------- डेटाबेस
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()


def get_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _pg():
    import psycopg2
    return psycopg2.connect(DATABASE_URL, sslmode="require")


def db_init():
    if DATABASE_URL:
        conn = _pg()
        try:
            with conn:
                with conn.cursor() as c:
                    c.execute("""
                        CREATE TABLE IF NOT EXISTS responses (
                            id SERIAL PRIMARY KEY,
                            created_at TEXT NOT NULL,
                            data TEXT NOT NULL
                        )
                    """)
        finally:
            conn.close()
    else:
        with get_db() as db:
            db.execute("""
                CREATE TABLE IF NOT EXISTS responses (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at TEXT NOT NULL,
                    data TEXT NOT NULL
                )
            """)


def db_insert(created_at, data_json):
    """नया उत्तर सहेजो और उसकी id लौटाओ (preview पेज हेतु)।"""
    if DATABASE_URL:
        conn = _pg()
        try:
            with conn:
                with conn.cursor() as c:
                    c.execute("INSERT INTO responses (created_at, data) VALUES (%s, %s) RETURNING id",
                              (created_at, data_json))
                    return c.fetchone()[0]
        finally:
            conn.close()
    else:
        with get_db() as db:
            cur = db.execute("INSERT INTO responses (created_at, data) VALUES (?, ?)",
                             (created_at, data_json))
            return cur.lastrowid


def db_get(rid):
    """एक उत्तर id से निकालो (preview पेज हेतु)।"""
    if DATABASE_URL:
        conn = _pg()
        try:
            with conn.cursor() as c:
                c.execute("SELECT id, created_at, data FROM responses WHERE id = %s", (rid,))
                row = c.fetchone()
                if not row:
                    return None
                cols = [d[0] for d in c.description]
                return dict(zip(cols, row))
        finally:
            conn.close()
    else:
        with get_db() as db:
            r = db.execute("SELECT id, created_at, data FROM responses WHERE id = ?",
                           (rid,)).fetchone()
            if not r:
                return None
            return {"id": r["id"], "created_at": r["created_at"], "data": r["data"]}


def db_all():
    """सभी उत्तर — नए पहले (id DESC)"""
    if DATABASE_URL:
        conn = _pg()
        try:
            with conn.cursor() as c:
                c.execute("SELECT id, created_at, data FROM responses ORDER BY id DESC")
                cols = [d[0] for d in c.description]
                return [dict(zip(cols, row)) for row in c.fetchall()]
        finally:
            conn.close()
    else:
        rows = []
        with get_db() as db:
            for r in db.execute("SELECT id, created_at, data FROM responses ORDER BY id DESC"):
                rows.append({"id": r["id"], "created_at": r["created_at"], "data": r["data"]})
        return rows


def db_clear():
    if DATABASE_URL:
        conn = _pg()
        try:
            with conn:
                with conn.cursor() as c:
                    c.execute("DELETE FROM responses")
        finally:
            conn.close()
    else:
        with get_db() as db:
            db.execute("DELETE FROM responses")


db_init()


def parse_data(row):
    try:
        return json.loads(row["data"])
    except Exception:
        return {}


# ---------------------------------------------------------------- सार्वजनिक फॉर्म
@app.route("/")
def index():
    return redirect(url_for("form_page"))


@app.route("/form")
def form_page():
    cur_month, _ = current_month_year()
    return render_template(
        "form.html",
        title=get_form_title(), desc=FORM_DESC, instruction=get_instruction(),
        sections=SECTIONS, part_subtitles=PART_SUBTITLES, current_month=cur_month,
        site=SITE_TITLE, org_name=ORG_NAME,
        org_tagline=ORG_TAGLINE, org_sub=ORG_SUB, logo=LOGO_EXISTS,
    )


@app.route("/submit", methods=["POST"])
def submit():
    # हनीपॉट — स्पैम बॉट यहाँ लिखते हैं
    if request.form.get("website"):
        return redirect(url_for("thank_you"))

    payload = {}
    for f in FLAT:
        payload[f["id"]] = (request.form.get(f["id"]) or "").strip()

    # योग वाले फ़ील्ड सर्वर पर दोबारा गिनो (सुरक्षा हेतु — browser चाहे जो भेजे)
    for f in FLAT:
        if f.get("calc"):
            total = 0
            for src in f["calc"]:
                total += to_num(payload.get(src)) or 0
            payload[f["id"]] = str(total)

    ok = True
    for f in FLAT:
        if not f["required"] or payload[f["id"]]:
            continue
        # सशर्त फ़ील्ड: शर्त पूरी होने पर ही ज़रूरी (जैसे "नहीं" चुनने पर ही तिथि)
        if f.get("show_if") and payload.get(f["show_if"]) != f.get("show_if_value"):
            continue
        ok = False

    if not ok:
        flash("कृपया सभी ज़रूरी (*) प्रश्न भरें।", "error")
        return redirect(url_for("form_page"))

    rid = db_insert(datetime.now().isoformat(timespec="seconds"),
                      json.dumps(payload, ensure_ascii=False))
    return redirect(url_for("preview", rid=rid))


@app.route("/preview/<int:rid>")
def preview(rid):
    """Submit के बाद preview — भरा हुआ पूरा विवरण, प्रिंट सहित।"""
    row = db_get(rid)
    if not row:
        abort(404)
    data = parse_data(row)
    return render_template(
        "preview.html", data=data, sections=SECTIONS, part_subtitles=PART_SUBTITLES,
        created=row["created_at"], rid=rid, msg=THANK_YOU_MSG,
        title=get_form_title(), site=SITE_TITLE,
        org_name=ORG_NAME, org_tagline=ORG_TAGLINE, org_sub=ORG_SUB, logo=LOGO_EXISTS,
    )


@app.route("/dhanyavaad")
def thank_you():
    return render_template("thank_you.html", msg=THANK_YOU_MSG, site=SITE_TITLE,
                           org_name=ORG_NAME, org_sub=ORG_SUB, logo=LOGO_EXISTS)


# ---------------------------------------------------------------- एडमिन प्रमाणीकरण
def make_admin_key(hours=24):
    ts = int(time.time()) + hours * 3600
    body = f"admin:{ts}".encode()
    sig = hmac.new(app.secret_key.encode(), body, hashlib.sha256).hexdigest()[:32]
    return body.decode() + "." + sig


def check_admin_key(key):
    if not key or "." not in key:
        return False
    body, sig = key.rsplit(".", 1)
    expected = hmac.new(app.secret_key.encode(), body.encode(), hashlib.sha256).hexdigest()[:32]
    if not hmac.compare_digest(sig, expected):
        return False
    try:
        parts = body.split(":")
        if parts[0] != "admin":
            return False
        if int(parts[1]) < time.time():
            return False
    except Exception:
        return False
    return True


def is_admin():
    """कुकी सत्र या URL कुंजी — दोनों में से कोई भी मान्य"""
    if session.get("admin"):
        return True
    k = request.args.get("key") or request.headers.get("X-Admin-Key")
    return bool(k) and check_admin_key(k)


@app.route("/admin/login", methods=["POST"])
def admin_login():
    pw = request.form.get("password", "")
    expected = os.environ.get("ADMIN_PASSWORD", ADMIN_PASSWORD)
    if hmac.compare_digest(pw, expected):
        session["admin"] = True
        return redirect(url_for("admin", key=make_admin_key()))
    flash("गलत पासवर्ड!", "error")
    return redirect(url_for("admin"))


@app.route("/admin/logout", methods=["POST"])
def admin_logout():
    session.pop("admin", None)
    return redirect(url_for("admin"))


@app.route("/admin/clear", methods=["POST"])
def admin_clear():
    if not session.get("admin"):
        abort(403)
    db_clear()
    return redirect(url_for("admin"))


# ---------------------------------------------------------------- एडमिन डैशबोर्ड
@app.route("/admin")
def admin():
    key = request.args.get("key")
    if key and check_admin_key(key):
        session["admin"] = True
    elif not session.get("admin"):
        return render_template("admin_login.html", site=SITE_TITLE,
                               org_name=ORG_NAME, logo=LOGO_EXISTS)
    export_url = url_for("admin_export", key=key) if key else url_for("admin_export")
    pdf1_url = url_for("admin_pdf1", key=key) if key else url_for("admin_pdf1")
    pdf2_url = url_for("admin_pdf2", key=key) if key else url_for("admin_pdf2")
    return render_template(
        "admin.html",
        title=get_form_title(), site=SITE_TITLE,
        export_url=export_url, pdf1_url=pdf1_url, pdf2_url=pdf2_url,
        org_name=ORG_NAME, org_tagline=ORG_TAGLINE, org_sub=ORG_SUB, logo=LOGO_EXISTS,
    )


def build_summary():
    """सभी उत्तरों से आँकड़े बनाओ"""
    rows = db_all()

    total = len(rows)
    today_str = date.today().isoformat()
    week_start = (date.today() - timedelta(days=6)).isoformat()
    today_count = sum(1 for r in rows if r["created_at"][:10] == today_str)
    week_count = sum(1 for r in rows if r["created_at"][:10] >= week_start)
    last_time = rows[0]["created_at"] if rows else None

    # पिछले 14 दिन की समय-श्रृंखला
    series = {}
    for i in range(13, -1, -1):
        d = (date.today() - timedelta(days=i)).strftime("%d/%m")
        series[d] = 0
    for r in rows:
        key = r["created_at"][8:10] + "/" + r["created_at"][5:7]
        if key in series:
            series[key] += 1

    # प्रति-फ़ील्ड योग (संख्यात्मक)
    field_sums = {}
    for f in FLAT:
        if f["type"] == "number":
            vals = [to_num(parse_data(r).get(f["id"])) for r in rows]
            vals = [v for v in vals if v is not None]
            field_sums[f["id"]] = sum(vals)
        else:
            field_sums[f["id"]] = None

    # ---------- उत्तर-सूची ट्रैकर: किस प्रांत ने भरा / नहीं भरा ----------
    def _norm(s):
        return (s or "").replace(" ", "").replace("\u3000", "").lower()

    filled = {}
    for r in rows:
        d = parse_data(r)
        prov = display_prov(d)
        if not prov:
            continue
        k = _norm(prov)
        if k in filled:
            continue
        filled[k] = {"name": prov, "at": r["created_at"],
                     "head": d.get("p1_head", "") or d.get("p2_coord", "")}

    tracker = {"expected": EXPECTED_PROVINCES, "status": [], "extra": []}
    for name in EXPECTED_PROVINCES:
        k = _norm(name)
        st = filled.pop(k, None)
        tracker["status"].append({
            "name": name, "filled": bool(st),
            "at": st["at"] if st else None, "head": st["head"] if st else None,
        })
    tracker["extra"] = [filled[k] for k in filled]
    tracker["filled"] = sum(1 for s in tracker["status"] if s["filled"])
    tracker["total"] = len(tracker["status"])
    tracker["pending"] = [s["name"] for s in tracker["status"] if not s["filled"]]

    return {
        "total": total,
        "today": today_count,
        "week": week_count,
        "last_time": last_time,
        "month_label": get_month_label(),
        "other_prov_label": OTHER_PROV_LABEL,
        "series": [{"label": k, "value": v} for k, v in series.items()],
        "fields": [
            {"id": f["id"], "label": f["label"], "short": f.get("short", ""),
             "type": f["type"], "part": f["part"], "part_short": PART_SHORT.get(f["part"], ""),
             "grp": f["grp"], "item_label": f["item_label"],
             "header": col_header(f),
             "show_if": f.get("show_if"), "show_if_value": f.get("show_if_value")}
            for f in FLAT
        ],
        "field_sums": field_sums,
        "tracker": tracker,
        "rows": [
            {"id": r["id"], "created_at": r["created_at"], "data": parse_data(r)}
            for r in rows
        ],
        "sections": [
            {"part": sec.get("part", ""), "title": sec["title"],
             "items": [{"grp": it.get("grp", ""), "label": it.get("label", "")}
                       for it in sec.get("items", [])]}
            for sec in SECTIONS
        ],
        "parts": [
            {"name": p, "subtitle": PART_SUBTITLES.get(p, "")} for p in PART_ORDER
        ],
    }


@app.route("/admin/api/data")
def admin_api_data():
    if not is_admin():
        return jsonify({"error": "unauthorized"}), 401
    return jsonify(build_summary())


# ---------------------------------------------------------------- Excel एक्सपोर्ट (single combined)
def build_workbook(s):
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ---- शीट 1: सारे उत्तर (प्रथम + द्वितीय भाग, एक ही table में)
    ws = wb.active
    ws.title = "उत्तर"
    headers = ["क्रमांक", "जमा करने का समय"] + [col_header(f) for f in FLAT]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for i, r in enumerate(s["rows"], start=1):
        ts = r["created_at"].replace("T", "  ")
        row_vals = [i, ts]
        for f in FLAT:
            raw = r["data"].get(f["id"], "")
            if f.get("show_if") and r["data"].get(f["show_if"]) != f.get("show_if_value"):
                row_vals.append("लागू नहीं")
            elif f["type"] == "number":
                n = to_num(raw)
                row_vals.append(n if n is not None else "")
            else:
                row_vals.append(raw)
        ws.append(row_vals)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 21
    col = 3
    for f in FLAT:
        ws.column_dimensions[get_column_letter(col)].width = min(36, 10 + len(f["label"]))
        col += 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"

    # ---- शीट 2: सारांश (भाग-वार कुल योग + प्रांत-वार)
    ws2 = wb.create_sheet("सारांश")
    ws2.append([get_form_title()])
    ws2["A1"].font = Font(bold=True, size=14, color="1D4ED8")
    ws2.append(["कुल उत्तर (प्रांत)", s["total"]])
    ws2.append(["आज के उत्तर", s["today"]])
    ws2.append(["पिछले 7 दिन", s["week"]])
    ws2.append(["अंतिम उत्तर", s["last_time"] or "—"])
    ws2.append([])

    for part in PART_ORDER:
        sub = PART_SUBTITLES.get(part, "")
        ws2.append([part + (f" — {sub}" if sub else "")])
        ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True, size=12, color="1D4ED8")
        for sec in s["sections"]:
            if sec["part"] != part:
                continue
            items = [it for it in sec["items"] if it["grp"]]
            if not items:
                continue
            ws2.append([sec["title"]])
            ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True)
            for it in items:
                ws2.append([f"{it['grp']} {it['label']}"])
                ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True)
                for f in FLAT:
                    if f["part"] == part and f["grp"] == it["grp"]:
                        v = s["field_sums"].get(f["id"])
                        txt = str(v) if v is not None else "—"
                        ws2.append(["", f"{f['label']} — {txt}"])
        ws2.append([])

    # प्रांत-वार एक नज़र (मुख्य आँकड़े)
    ws2.append(["प्रांत-वार (मुख्य आँकड़े)"])
    ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True, color="1D4ED8")
    key_fields = ["p1_prov", "p1_head", "p1_kha1_cur", "p1_kha1_tgt",
                  "p1_g1_tot", "p1_g2_tot", "p1_d1_tgt", "p1_d2_done",
                  "p2_g1_tot", "p2_g2_tot"]
    ws2.append(["प्रांत", "अध्यक्ष/प्रमुख", "जिले वर्तमान", "जिले लक्ष्य",
                "दैनिक योग", "मासिक योग", "यज्ञ लक्ष्य", "यज्ञ सम्पन्न",
                "केन्द्र योग", "छात्र योग"])
    for c in range(1, len(key_fields) + 1):
        ws2.cell(row=ws2.max_row, column=c).font = Font(bold=True)
    for r in s["rows"]:
        vals = []
        for k in key_fields:
            if k == "p1_prov":
                vals.append(display_prov(r["data"]))
            else:
                vals.append(r["data"].get(k, ""))
        ws2.append(vals)

    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["B"].width = 24
    for c in range(3, len(key_fields) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    return wb


@app.route("/admin/export")
def admin_export():
    """सीधा Excel डाउनलोड — single combined"""
    if not is_admin():
        abort(403)
    wb = build_workbook(build_summary())
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"kary-vritt_{date.today().isoformat()}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------- PDF एक्सपोर्ट (भाग-वार, सिर्फ table)
def _pdf_groups(part):
    """भाग के फ़ील्ड → (section, grp) क्रम में समूह। परिचय (grp='') एक table में।"""
    groups, seen = [], {}
    for f in FLAT:
        if f["part"] != part:
            continue
        key = (f["section"], f["grp"])
        if key not in seen:
            seen[key] = {"section": f["section"], "grp": f["grp"],
                         "label": f["item_label"], "fields": []}
            groups.append(seen[key])
        if f["grp"] == "":
            seen[key]["label"] = f["section"]  # परिचय table का शीर्षक
        seen[key]["fields"].append(f)
    return groups


def build_pdf_bytes(part):
    """एक भाग की PDF — हर उप-खंड की table (पहला column: प्रांत)।"""
    reg = os.path.join(FONTS_DIR, "NotoSansDevanagari-Regular.ttf")
    bold = os.path.join(FONTS_DIR, "NotoSansDevanagari-Bold.ttf")
    if not (os.path.exists(reg) and os.path.exists(bold)):
        raise RuntimeError("Hindi font files missing in static/fonts/")

    s = build_summary()

    pdf = FPDF(orientation="L", format="A4")
    pdf.set_auto_page_break(True, margin=14)
    pdf.set_margins(10, 12, 10)
    pdf.add_font("Noto", "", reg)
    pdf.add_font("Noto", "B", bold)
    try:
        pdf.set_text_shaping(True)  # मात्राएँ/संयुक्ताक्षर सही दिखें
    except Exception as e:
        print("PDF shaping unavailable:", e)

    pdf.add_page()
    pdf.set_font("Noto", "B", 15)
    pdf.cell(0, 9, get_form_title(), new_x="LMARGIN", new_y="NEXT", align="C")
    sub = PART_SUBTITLES.get(part, "")
    head2 = part + (f" — {sub}" if sub else "") + f"   •   कुल उत्तर: {s['total']}"
    pdf.set_font("Noto", "", 11)
    pdf.cell(0, 7, head2, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(3)

    if not s["rows"]:
        pdf.set_font("Noto", "", 12)
        pdf.cell(0, 10, "अभी तक कोई उत्तर नहीं आया।", align="C")
        return bytes(pdf.output())

    usable = pdf.w - pdf.l_margin - pdf.r_margin
    head_style = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(29, 78, 216))

    for g in _pdf_groups(part):
        is_intro = (g["grp"] == "")
        title = g["label"] if is_intro else f"{g['grp']} · {g['label']}"
        pdf.set_font("Noto", "B", 11)
        pdf.cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")

        cols = g["fields"]
        # प्रथम भाग की परिचय-table में प्रांत खुद है; बाकी सब में प्रांत column जोड़ो
        has_own_prov = any(f["id"] == "p1_prov" for f in cols)
        if is_intro and has_own_prov:
            headers = ["क्र."] + [(f["short"] or f["label"]) for f in cols]
            widths = [10] + [max(30, (usable - 10) / max(len(cols), 1))] * len(cols)
        else:
            headers = ["क्र.", "प्रांत"] + [(f["short"] or f["label"]) for f in cols]
            rest = (usable - 10 - 34) / max(len(cols), 1)
            widths = [10, 34] + [max(22, rest)] * len(cols)
        # चौड़ाई को पेज में समायोजित करो
        scale = usable / sum(widths)
        widths = [w * scale for w in widths]

        pdf.set_font("Noto", "", 8.5)
        with pdf.table(col_widths=tuple(widths), text_align="LEFT",
                       line_height=5.2, width=usable,
                       headings_style=head_style, repeat_headings=1) as table:
            hdr = table.row()
            for h in headers:
                hdr.cell(h)
            for i, r in enumerate(s["rows"], start=1):
                row = table.row()
                row.cell(str(i), align="CENTER")
                if not (is_intro and has_own_prov):
                    row.cell(display_prov(r["data"]) or "—")
                for f in cols:
                    if is_intro and f["id"] == "p1_prov":
                        row.cell(r["data"].get(f["id"], "") or "—")
                        continue
                    if f.get("show_if") and r["data"].get(f["show_if"]) != f.get("show_if_value"):
                        row.cell("लागू नहीं")
                    elif f["type"] == "number":
                        n = to_num(r["data"].get(f["id"], ""))
                        row.cell("" if n is None else str(n), align="CENTER")
                    else:
                        row.cell(r["data"].get(f["id"], "") or "—")
        pdf.ln(4)

    return bytes(pdf.output())


@app.route("/admin/pdf1")
def admin_pdf1():
    """PDF-1 : प्रथम भाग — सभी उत्तरों की table"""
    if not is_admin():
        abort(403)
    data = build_pdf_bytes("प्रथम भाग")
    fname = f"kary-vritt_bhag-1_{date.today().isoformat()}.pdf"
    return send_file(io.BytesIO(data), as_attachment=True, download_name=fname,
                     mimetype="application/pdf")


@app.route("/admin/pdf2")
def admin_pdf2():
    """PDF-2 : द्वितीय भाग — सभी उत्तरों की table"""
    if not is_admin():
        abort(403)
    data = build_pdf_bytes("द्वितीय भाग")
    fname = f"kary-vritt_bhag-2_{date.today().isoformat()}.pdf"
    return send_file(io.BytesIO(data), as_attachment=True, download_name=fname,
                     mimetype="application/pdf")


# ---------------------------------------------------------------- रन
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    app.run(host="0.0.0.0", port=port, debug=False)
